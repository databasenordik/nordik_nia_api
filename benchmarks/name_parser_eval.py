"""Score the name parser against the hand-curated gold set.

``final_list_master.csv`` is the gold set: Master-list rows where a person decided, cell by
cell, which fact each fragment of a name represents. It supersedes
``master_name_format_lists.xlsx`` and adds curated ``_indigenous`` columns, so the parser's
Indigenous-name classification is now scored rather than taken on trust. This harness is
the gate for the parser -- it is a test fixture only, exactly as
``docs/supplied_questions_pipeline_results.md`` is for the question suite. No value from it
is compiled into ``backend/app``.

Either format is accepted, chosen by suffix, so an older workbook still runs.

Usage::

    python -m benchmarks.name_parser_eval [path/to/final_list_master.csv]
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.normalization.name_parser import NameParts, parse_master  # noqa: E402

DEFAULT_WORKBOOK = BACKEND_ROOT.parent / "final_list_master.csv"

# curated column -> attribute on NameParts
SCALAR_COLUMNS = {"_first": "first", "_last": "last", "_indigenous": "indigenous"}
LIST_COLUMNS = {
    "_middle": "middle",
    "_extra_first": "extra_first",
    "_extra_last": "extra_last",
    "_first_spelling": "first_spelling",
    "_last_spelling": "last_spelling",
    "_first_meaning": "first_meaning",
    "_last_meaning": "last_meaning",
    "_indigenous_spelling": "indigenous_spelling",
    "_indigenous_meaning": "indigenous_meaning",
}


def _cells(value: object) -> set[str]:
    if value is None:
        return set()
    return {part.strip() for part in str(value).split(";") if part.strip()}


def load_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        return _load_csv(path)
    return _load_workbook(path)


def _load_csv(path: Path) -> list[dict[str, str]]:
    """Read the curated CSV, keeping only the columns it actually names.

    The export carries trailing unnamed columns; they hold no data and would otherwise
    become dict keys of ''. utf-8-sig because Excel writes a BOM.
    """
    import csv

    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = [str(name).strip() for name in next(reader)]
        keep = [(index, name) for index, name in enumerate(header) if name]
        rows: list[dict[str, str]] = []
        for raw in reader:
            record = {
                name: (raw[index].strip() if index < len(raw) and raw[index] else "")
                for index, name in keep
            }
            if any(record.values()):
                record["__sheet__"] = path.stem
                rows.append(record)
    return rows


def _load_workbook(path: Path) -> list[dict[str, str]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict[str, str]] = []
    for sheet in workbook.worksheets:
        header = [cell.value for cell in sheet[1]][:13]
        for raw in sheet.iter_rows(min_row=2, max_col=13, values_only=True):
            record = {
                str(key): ("" if value is None else str(value).strip())
                for key, value in zip(header, raw, strict=False)
                if key is not None
            }
            if any(record.values()):
                record["__sheet__"] = sheet.title
                rows.append(record)
    return rows


def _superseded(source: str, variant: str) -> str:
    """Whether curation stored something the parser now deliberately stores differently.

    Three categories were reviewed against the raw source and reclassified. The curated
    workbook predates those decisions, so counting them as losses would hold the gate to a
    standard the code intentionally no longer follows. They are reported separately rather
    than silently dropped, so the divergence stays visible.
    """
    bare = variant.strip().rstrip(".").lower()
    if bare in {"sr", "jr", "snr", "jnr", "ii", "iii", "iv"}:
        # A generational suffix identifies a person but is not an alternate given name;
        # it is now carried as a flag.
        return "generational suffix -> flag"
    if variant.strip().startswith("["):
        # Whole-name brackets are the transcriber's doubt about the reading, not a second
        # spelling; the unbracketed form is canonical and the cell is flagged uncertain.
        return "whole-name bracket -> uncertain flag"
    if variant.strip().endswith("?") or "(sp?)" in source or "(?)" in source:
        # A question mark is uncertainty about the name, not a variant of it.
        return "uncertainty marker -> uncertain flag"
    return ""


def _classification(rows: list[dict[str, str]]) -> tuple[int, int, int, int, Counter[str]]:
    """Return (correct, misplaced, never_extracted, total, superseded) over the corpus."""
    correct = misplaced = unextracted = total = 0
    superseded: Counter[str] = Counter()
    for row in rows:
        parts = parse_master(
            first=row.get("first", ""), middle=row.get("middle", ""), last=row.get("last", "")
        )
        anywhere = set(parts.searchable())
        for part in ("first", "last"):
            want_extra = _cells(row.get(f"_extra_{part}", ""))
            want_spelling = _cells(row.get(f"_{part}_spelling", ""))
            got_extra = set(getattr(parts, f"extra_{part}"))
            got_spelling = set(getattr(parts, f"{part}_spelling"))
            source = row.get(part, "")
            for value in want_extra | want_spelling:
                if value in got_extra | got_spelling:
                    total += 1
                    if (value in want_extra) == (value in got_extra):
                        correct += 1
                    else:
                        misplaced += 1
                elif value not in anywhere:
                    reason = _superseded(source, value)
                    if reason:
                        superseded[reason] += 1
                    else:
                        total += 1
                        unextracted += 1
                else:
                    total += 1
    return correct, misplaced, unextracted, total, superseded


def evaluate(rows: list[dict[str, str]]) -> int:
    scalar_hits: Counter[str] = Counter()
    scalar_total: Counter[str] = Counter()
    list_tp: Counter[str] = Counter()
    list_fp: Counter[str] = Counter()
    list_fn: Counter[str] = Counter()
    rescued: Counter[str] = Counter()
    misses: list[tuple[str, str, str, str, str]] = []

    for row in rows:
        parts: NameParts = parse_master(
            first=row.get("first", ""),
            middle=row.get("middle", ""),
            last=row.get("last", ""),
        )
        alias = set(parts.alias)

        for column, attribute in SCALAR_COLUMNS.items():
            expected = row.get(column, "").strip()
            if not expected:
                continue
            actual = getattr(parts, attribute)
            scalar_total[column] += 1
            if actual == expected:
                scalar_hits[column] += 1
            elif len(misses) < 400:
                misses.append((column, row.get("first", ""), row.get("last", ""), expected, actual))

        for column, attribute in LIST_COLUMNS.items():
            expected = _cells(row.get(column, ""))
            actual = set(getattr(parts, attribute))
            if not expected and not actual:
                continue
            list_tp[column] += len(expected & actual)
            missing = expected - actual
            list_fn[column] += len(missing)
            list_fp[column] += len(actual - expected)
            # A variant the rules could not classify is meant to land in alias. It is
            # still findable there, so count it separately from a genuine loss.
            rescued[column] += len(missing & alias)

    print(f"rows evaluated: {len(rows)}\n")
    # _first used to be gated below _last because the old workbook was itself undecided
    # about slashed spellings -- five cells chose the first form, five the second.
    # final_list_master.csv settles it: the first-written form, seven times out of seven.
    # With the parser following that, all three components sit at 99.9% or better, so the
    # gates rise to match. Two residual misses are known and accepted: "Garden Ri" is a
    # two-word surname the parser splits, and "Mona." keeps a trailing period that the
    # general rule strips.
    gates = {"_first": 0.99, "_last": 0.99, "_indigenous": 0.99}
    print("EXACT COMPONENT (canonical first / last)")
    ok = True
    for column in SCALAR_COLUMNS:
        total = scalar_total[column]
        if not total:
            continue
        acc = scalar_hits[column] / total
        gate = gates[column]
        flag = "OK " if acc >= gate else "XX "
        ok = ok and acc >= gate
        print(
            f"  {flag}{column:16} {scalar_hits[column]:>4}/{total:<4} = {acc:6.1%}"
            f"   (gate {gate:.0%})"
        )

    print("\nVARIANT COLUMNS (set overlap per cell)")
    grand_tp = grand_fp = grand_fn = grand_rescued = 0
    for column in LIST_COLUMNS:
        tp, fp, fn = list_tp[column], list_fp[column], list_fn[column]
        if not (tp or fp or fn):
            continue
        grand_tp += tp
        grand_fp += fp
        grand_fn += fn
        grand_rescued += rescued[column]
        precision = tp / (tp + fp) if tp + fp else 1.0
        recall = tp / (tp + fn) if tp + fn else 1.0
        print(
            f"     {column:20} P={precision:6.1%} R={recall:6.1%} "
            f"(tp={tp} fp={fp} fn={fn}, {rescued[column]} of the misses landed in alias)"
        )

    denominator = grand_tp + grand_fn
    exact = grand_tp / denominator if denominator else 1.0
    findable = (grand_tp + grand_rescued) / denominator if denominator else 1.0

    # Extraction and classification are separate failures and are gated separately. A
    # variant that was never pulled out of the cell cannot be found at all; a variant in
    # the wrong column is still found. Rolling them into one number hides which is broken,
    # and is not comparable to the 93.7% baseline, which measured classification alone.
    correct, misplaced, unextracted, total, superseded = _classification(rows)
    classification = correct / (correct + misplaced) if correct + misplaced else 1.0
    recall = (total - unextracted) / total if total else 1.0
    print(
        f"\n  variant placement  : {exact:.1%} exact, {findable:.1%} findable counting alias"
        f"\n  classification     : {classification:.1%}  (gate 93.7%, the curated baseline)"
        f"\n  extraction recall  : {recall:.1%}  (gate 95%) -- {unextracted} never extracted"
    )
    if superseded:
        print(
            f"\n  excluded: {sum(superseded.values())} curated variants the parser now"
            " stores differently by decision, not by loss"
        )
        for reason, count in superseded.most_common():
            print(f"      {count:>3}  {reason}")
    ok = ok and classification >= 0.937 and recall >= 0.95

    if misses:
        print("\nCOMPONENT MISSES (first 15)")
        for column, first, last, expected, actual in misses[:15]:
            source = first if "first" in column else last
            print(f"  {column:16} src={source[:40]!r:42} want={expected!r:20} got={actual!r}")

    print("\nRESULT:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    if not path.exists():
        raise SystemExit(f"gold set not found: {path}")
    return evaluate(load_rows(path))


if __name__ == "__main__":
    raise SystemExit(main())
